from __future__ import annotations

import io
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, List, Optional, Set, Tuple

import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from PIL import Image
from tqdm import tqdm

from models.data_models import ExtractionResult, ImageBlock, SheetDependency, TableBlock
from utils.helpers import (
    dedupe_key_from_pil,
    detect_csv_delimiter,
    detect_csv_encoding,
    ensure_dir,
    setup_logger,
)

FORMULA_CELL_PATTERN = re.compile(
    r"(?:(?:'[^']+?'|[A-Za-z0-9_]+)!)?\$?[A-Z]+\$?[0-9]+(?::\$?[A-Z]+\$?[0-9]+)?"
)
SHEET_REF_PATTERN = re.compile(r"(?:'([^']+)'|([A-Za-z0-9_]+))!")


def _is_numeric(value: str) -> bool:
    value = (value or "").strip()
    if not value:
        return False
    return bool(re.fullmatch(r"[-+]?\d*\.?\d+(?:[eE][-+]?\d+)?", value))


def _detect_headers(rows: List[List[str]]) -> Optional[List[str]]:
    for row in rows[:10]:
        if not row:
            continue
        non_numeric = sum(1 for cell in row if cell.strip() and not _is_numeric(cell))
        if non_numeric >= max(1, len(row) // 2):
            return row
    return rows[0] if rows else None


def _sheet_matrix(ws: Worksheet) -> Tuple[List[List[str]], bool]:
    max_row = ws.max_row or 0
    max_col = ws.max_column or 0
    matrix = [["" for _ in range(max_col)] for _ in range(max_row)]

    for row in ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
        for cell in row:
            value = "" if cell.value is None else str(cell.value)
            matrix[cell.row - 1][cell.column - 1] = value

    has_merged = bool(ws.merged_cells.ranges)
    for merge_range in ws.merged_cells.ranges:
        min_col, min_row, max_col, max_row = merge_range.bounds
        anchor_val = matrix[min_row - 1][min_col - 1]
        for r in range(min_row - 1, max_row):
            for c in range(min_col - 1, max_col):
                matrix[r][c] = anchor_val
    return matrix, has_merged


def _formula_dependencies(sheet_name: str, ws: Worksheet, workbook_sheet_names: Set[str]) -> SheetDependency:
    references: Set[str] = set()
    formulas_with_cross_refs: Dict[str, str] = {}

    for row in ws.iter_rows():
        for cell in row:
            value = cell.value
            if not isinstance(value, str) or not value.startswith("="):
                continue

            formula = value
            cell_ref = f"{cell.coordinate}"
            sheet_matches = SHEET_REF_PATTERN.findall(formula)
            targets = {m[0] or m[1] for m in sheet_matches if (m[0] or m[1])}
            valid_targets = {t for t in targets if t in workbook_sheet_names and t != sheet_name}
            if valid_targets:
                references.update(valid_targets)
                formulas_with_cross_refs[cell_ref] = formula
            elif FORMULA_CELL_PATTERN.search(formula):
                # Keep formula mapping even if same-sheet reference for transparency.
                formulas_with_cross_refs[cell_ref] = formula

    return SheetDependency(
        sheet_name=sheet_name,
        references=sorted(references),
        formulas_with_cross_refs=formulas_with_cross_refs,
    )


def _extract_worksheet_images(
    ws: Worksheet,
    sheet_idx: int,
    images_dir: Path,
    dedupe_hashes: set[str],
) -> List[ImageBlock]:
    blocks: List[ImageBlock] = []
    img_counter = 0
    for image_obj in getattr(ws, "_images", []):
        try:
            data = image_obj._data()
            pil = Image.open(io.BytesIO(data)).convert("RGB")
        except Exception:
            continue

        dedupe_key = dedupe_key_from_pil(pil)
        if dedupe_key in dedupe_hashes:
            continue
        dedupe_hashes.add(dedupe_key)

        img_counter += 1
        out_name = f"sheet{sheet_idx}_img{img_counter}.png"
        out_path = images_dir / out_name
        pil.save(out_path)
        blocks.append(
            ImageBlock(
                image_path=str(out_path),
                page_or_slide=sheet_idx,
                bounding_box=None,
                dimensions={"width": pil.width, "height": pil.height},
                extraction_method="embedded",
            )
        )
    return blocks


def _extract_csv(file_path: Path) -> Tuple[List[TableBlock], List[str]]:
    warnings: List[str] = []
    encoding = detect_csv_encoding(file_path)
    delimiter = detect_csv_delimiter(file_path, encoding)

    try:
        df = pd.read_csv(
            file_path,
            delimiter=delimiter,
            encoding=encoding,
            engine="python",
            dtype=str,
            keep_default_na=False,
        )
    except Exception as exc:
        warnings.append(f"CSV parsing failed with detected options; fallback attempted: {exc}")
        df = pd.read_csv(file_path, engine="python", dtype=str, keep_default_na=False)

    rows = [[str(v) if v is not None else "" for v in row] for row in df.fillna("").values.tolist()]
    header = list(df.columns.astype(str))
    if rows and all((h.startswith("Unnamed:") for h in header)):
        header = _detect_headers(rows)

    table = TableBlock(
        data=[header] + rows if header else rows,
        headers=header,
        page_or_slide=1,
        sheet_name=file_path.stem,
        num_rows=len(rows) + (1 if header else 0),
        num_cols=len(header) if header else (max((len(r) for r in rows), default=0)),
        has_merged_cells=False,
        source_formulas=None,
    )
    return [table], warnings


def _extract_xlsx(file_path: Path, images_dir: Path) -> Tuple[List[TableBlock], List[SheetDependency], List[ImageBlock], List[str]]:
    warnings: List[str] = []
    tables: List[TableBlock] = []
    dependencies: List[SheetDependency] = []
    images: List[ImageBlock] = []
    dedupe_hashes: set[str] = set()

    wb = load_workbook(file_path, data_only=False, read_only=False)
    sheet_names_set = set(wb.sheetnames)
    named_ranges = list(getattr(wb, "defined_names", []))

    for sheet_idx, ws in enumerate(tqdm(wb.worksheets, desc=f"Extracting workbook {file_path.name}"), start=1):
        matrix, has_merged = _sheet_matrix(ws)
        headers = _detect_headers(matrix)
        formula_map: Dict[str, str] = {}
        dep = _formula_dependencies(ws.title, ws, sheet_names_set)
        dependencies.append(dep)
        formula_map.update(dep.formulas_with_cross_refs)

        metadata_formulas = dict(formula_map)
        metadata_formulas["__charts__"] = str(len(getattr(ws, "_charts", [])))
        metadata_formulas["__conditional_formatting_rules__"] = str(len(ws.conditional_formatting))
        data_validations = getattr(ws.data_validations, "dataValidation", []) if ws.data_validations else []
        metadata_formulas["__data_validations__"] = str(len(data_validations))
        metadata_formulas["__named_ranges__"] = ",".join(str(n) for n in named_ranges)

        tables.append(
            TableBlock(
                data=matrix,
                headers=headers,
                page_or_slide=sheet_idx,
                sheet_name=ws.title,
                num_rows=len(matrix),
                num_cols=max((len(r) for r in matrix), default=0),
                has_merged_cells=has_merged,
                source_formulas=metadata_formulas if metadata_formulas else None,
            )
        )

        images.extend(_extract_worksheet_images(ws, sheet_idx, images_dir, dedupe_hashes))

    wb.close()
    return tables, dependencies, images, warnings


def extract_spreadsheet(file_path: str | Path, output_dir: str | Path) -> ExtractionResult:
    logger = setup_logger()
    file_path = Path(file_path)
    base_output = ensure_dir(output_dir)
    images_dir = ensure_dir(base_output / "images")
    file_type = file_path.suffix.lower().lstrip(".")

    tables: List[TableBlock] = []
    images: List[ImageBlock] = []
    warnings: List[str] = []
    dependencies: Optional[List[SheetDependency]] = None

    try:
        if file_type == "csv":
            tables, csv_warnings = _extract_csv(file_path)
            warnings.extend(csv_warnings)
            num_units = 1
        else:
            tables, dependencies, images, sheet_warnings = _extract_xlsx(file_path, images_dir)
            warnings.extend(sheet_warnings)
            num_units = len({tb.sheet_name for tb in tables if tb.sheet_name})
    except Exception as exc:
        warnings.append(f"Spreadsheet extraction failed: {exc}")
        logger.warning(warnings[-1])
        num_units = 0

    return ExtractionResult(
        file_name=file_path.name,
        file_type=file_type,
        extraction_timestamp=datetime.now(timezone.utc),
        num_pages_or_slides=max(num_units, 1),
        texts=[],
        tables=tables,
        images=images,
        sheet_dependencies=dependencies,
        warnings=warnings,
    )
