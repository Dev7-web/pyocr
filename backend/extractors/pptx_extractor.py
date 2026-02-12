from __future__ import annotations

import io
import subprocess
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

import fitz
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from tqdm import tqdm

from models.data_models import ExtractionResult, ImageBlock, TableBlock, TextBlock
from utils.helpers import dedupe_key_from_pil, ensure_dir, setup_logger


def _iter_shapes_recursive(shapes: Sequence) -> Iterable:
    for shape in shapes:
        yield shape
        if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes_recursive(shape.shapes)


def _detect_header_row(rows: List[List[str]]) -> Optional[List[str]]:
    if not rows:
        return None
    first = rows[0]
    if not first:
        return None
    alpha_like = sum(1 for c in first if any(ch.isalpha() for ch in c))
    return first if alpha_like >= max(1, len(first) // 2) else None


def _extract_texts_and_tables(prs: Presentation) -> Tuple[List[TextBlock], List[TableBlock], List[int]]:
    texts: List[TextBlock] = []
    tables: List[TableBlock] = []
    render_needed_slides: List[int] = []
    diagram_type = getattr(MSO_SHAPE_TYPE, "DIAGRAM", None)
    chart_type = getattr(MSO_SHAPE_TYPE, "CHART", None)

    for slide_idx, slide in enumerate(prs.slides, start=1):
        has_complex_visual = False
        for shape in _iter_shapes_recursive(slide.shapes):
            try:
                if getattr(shape, "has_text_frame", False) and shape.text_frame is not None:
                    content = shape.text_frame.text.strip()
                    if content:
                        texts.append(
                            TextBlock(
                                content=content,
                                page_or_slide=slide_idx,
                                block_type="paragraph",
                            )
                        )

                if getattr(shape, "shape_type", None) in {chart_type, diagram_type}:
                    has_complex_visual = True
                if getattr(shape, "has_chart", False):
                    has_complex_visual = True

                if getattr(shape, "has_table", False):
                    table = shape.table
                    rows: List[List[str]] = []
                    for row in table.rows:
                        rows.append([cell.text_frame.text.strip() if cell.text_frame else "" for cell in row.cells])
                    if rows:
                        tables.append(
                            TableBlock(
                                data=rows,
                                headers=_detect_header_row(rows),
                                page_or_slide=slide_idx,
                                num_rows=len(rows),
                                num_cols=max((len(r) for r in rows), default=0),
                            )
                        )
            except Exception:
                continue

        notes = getattr(slide, "notes_slide", None)
        if notes and notes.notes_text_frame:
            note_text = notes.notes_text_frame.text.strip()
            if note_text:
                texts.append(TextBlock(content=note_text, page_or_slide=slide_idx, block_type="note"))

        if has_complex_visual:
            render_needed_slides.append(slide_idx)

    return texts, tables, render_needed_slides


def _save_picture_shape(shape, slide_idx: int, img_counter: int, images_dir: Path) -> Optional[Tuple[ImageBlock, int, str]]:
    try:
        blob = shape.image.blob
        pil = Image.open(io.BytesIO(blob)).convert("RGB")
        ext = (shape.image.ext or "png").lower()
        img_counter += 1
        image_name = f"slide{slide_idx}_img{img_counter}.{ext}"
        image_path = images_dir / image_name
        pil.save(image_path)

        bbox = None
        if hasattr(shape, "left") and hasattr(shape, "top"):
            bbox = {
                "x0": float(shape.left),
                "y0": float(shape.top),
                "x1": float(shape.left + shape.width),
                "y1": float(shape.top + shape.height),
            }

        block = ImageBlock(
            image_path=str(image_path),
            page_or_slide=slide_idx,
            bounding_box=bbox,
            dimensions={"width": pil.width, "height": pil.height},
            extraction_method="embedded",
        )
        return block, img_counter, dedupe_key_from_pil(pil)
    except Exception:
        return None


def _extract_embedded_images(prs: Presentation, images_dir: Path) -> List[ImageBlock]:
    image_blocks: List[ImageBlock] = []
    dedupe_hashes: set[str] = set()
    img_counter = 0
    picture_type = getattr(MSO_SHAPE_TYPE, "PICTURE", None)

    for slide_idx, slide in enumerate(prs.slides, start=1):
        for shape in _iter_shapes_recursive(slide.shapes):
            if getattr(shape, "shape_type", None) != picture_type:
                continue
            payload = _save_picture_shape(shape, slide_idx, img_counter, images_dir)
            if payload is None:
                continue
            block, img_counter, dedupe_key = payload
            if dedupe_key in dedupe_hashes:
                continue
            dedupe_hashes.add(dedupe_key)
            image_blocks.append(block)

    return image_blocks


def _render_slides_via_libreoffice(
    file_path: Path,
    slides_to_render: Sequence[int],
    images_dir: Path,
) -> List[ImageBlock]:
    if not slides_to_render:
        return []

    logger = setup_logger()
    rendered_blocks: List[ImageBlock] = []

    with tempfile.TemporaryDirectory(prefix="pptx_render_") as tmp_dir:
        tmp_path = Path(tmp_dir)
        cmd = [
            "libreoffice",
            "--headless",
            "--convert-to",
            "pdf",
            str(file_path),
            "--outdir",
            str(tmp_path),
        ]
        try:
            subprocess.run(cmd, check=True, capture_output=True, text=True)
        except Exception as exc:
            logger.warning("LibreOffice conversion failed for %s: %s", file_path.name, exc)
            return []

        pdf_path = tmp_path / f"{file_path.stem}.pdf"
        if not pdf_path.exists():
            matches = list(tmp_path.glob("*.pdf"))
            if not matches:
                return []
            pdf_path = matches[0]

        try:
            doc = fitz.open(pdf_path)
        except Exception as exc:
            logger.warning("Failed to open rendered PDF for %s: %s", file_path.name, exc)
            return []

        chart_counter = 0
        target_set = set(slides_to_render)
        for page_idx in range(doc.page_count):
            slide_num = page_idx + 1
            if slide_num not in target_set:
                continue
            page = doc[page_idx]
            pix = page.get_pixmap(dpi=220, alpha=False)
            pil = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            chart_counter += 1
            out_name = f"slide{slide_num}_chart{chart_counter}.png"
            out_path = images_dir / out_name
            pil.save(out_path)
            rendered_blocks.append(
                ImageBlock(
                    image_path=str(out_path),
                    page_or_slide=slide_num,
                    bounding_box=None,
                    dimensions={"width": pil.width, "height": pil.height},
                    extraction_method="slide_render",
                )
            )

        doc.close()

    return rendered_blocks


def extract_pptx(file_path: str | Path, output_dir: str | Path) -> ExtractionResult:
    logger = setup_logger()
    file_path = Path(file_path)
    base_output = ensure_dir(output_dir)
    images_dir = ensure_dir(base_output / "images")

    try:
        prs = Presentation(str(file_path))
    except Exception as exc:
        raise RuntimeError(f"Failed to open PPTX: {file_path}") from exc

    warnings: List[str] = []
    texts: List[TextBlock] = []
    tables: List[TableBlock] = []
    images: List[ImageBlock] = []

    try:
        texts, tables, render_slides = _extract_texts_and_tables(prs)
    except Exception as exc:
        warnings.append(f"Failed PPTX text/table extraction: {exc}")
        logger.warning(warnings[-1])
        render_slides = []

    try:
        images.extend(_extract_embedded_images(prs, images_dir))
    except Exception as exc:
        warnings.append(f"Failed PPTX image extraction: {exc}")
        logger.warning(warnings[-1])

    try:
        for block in tqdm(
            _render_slides_via_libreoffice(file_path, render_slides, images_dir),
            desc=f"Rendering charts {file_path.name}",
        ):
            images.append(block)
    except Exception as exc:
        warnings.append(f"Failed PPTX slide render fallback: {exc}")
        logger.warning(warnings[-1])

    return ExtractionResult(
        file_name=file_path.name,
        file_type="pptx",
        extraction_timestamp=datetime.now(timezone.utc),
        num_pages_or_slides=len(prs.slides),
        texts=texts,
        tables=tables,
        images=images,
        warnings=warnings,
    )
