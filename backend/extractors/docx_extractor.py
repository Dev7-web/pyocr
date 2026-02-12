from __future__ import annotations

import hashlib
import io
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable, List, Optional, Tuple

from docx import Document
from docx.table import Table
from PIL import Image
from tqdm import tqdm

from models.data_models import ExtractionResult, ImageBlock, TableBlock, TextBlock
from utils.helpers import dedupe_key_from_pil, ensure_dir, setup_logger


def _heading_level(style_name: str | None) -> Optional[int]:
    if not style_name:
        return None
    match = re.match(r"Heading\s+(\d+)", style_name, flags=re.IGNORECASE)
    if not match:
        return None
    return int(match.group(1))


def _extract_paragraph_texts(doc: Document) -> List[TextBlock]:
    blocks: List[TextBlock] = []
    current_section: Optional[str] = None

    for idx, para in enumerate(doc.paragraphs, start=1):
        text = (para.text or "").strip()
        if not text:
            continue
        level = _heading_level(getattr(para.style, "name", None))
        if level is not None:
            current_section = text
            block_type = "heading"
        else:
            block_type = "paragraph"
        blocks.append(
            TextBlock(
                content=text,
                page_or_slide=idx,
                heading_level=level,
                source_section=current_section,
                block_type=block_type,
            )
        )
    return blocks


def _extract_header_footer_texts(doc: Document, start_index: int) -> List[TextBlock]:
    blocks: List[TextBlock] = []
    index = start_index
    seen = set()

    for section in doc.sections:
        for area_name, area in [("header", section.header), ("footer", section.footer)]:
            cache_key = (id(area), area_name)
            if cache_key in seen:
                continue
            seen.add(cache_key)
            for para in area.paragraphs:
                text = (para.text or "").strip()
                if not text:
                    continue
                blocks.append(TextBlock(content=text, page_or_slide=index, block_type=area_name))
                index += 1
    return blocks


def _table_has_merge(table: Table) -> bool:
    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.tcPr
            if tc_pr is None:
                continue
            if tc_pr.gridSpan is not None:
                return True
            if tc_pr.vMerge is not None:
                return True
    return False


def _extract_table_recursive(table: Table, page_or_slide: int, out: List[TableBlock]) -> None:
    rows: List[List[str]] = []
    for row in table.rows:
        rows.append([cell.text.strip() for cell in row.cells])
    if rows:
        out.append(
            TableBlock(
                data=rows,
                headers=rows[0] if rows else None,
                page_or_slide=page_or_slide,
                num_rows=len(rows),
                num_cols=max((len(r) for r in rows), default=0),
                has_merged_cells=_table_has_merge(table),
            )
        )

    for row in table.rows:
        for cell in row.cells:
            for nested in cell.tables:
                _extract_table_recursive(nested, page_or_slide, out)


def _extract_tables(doc: Document) -> List[TableBlock]:
    out: List[TableBlock] = []
    for idx, table in enumerate(doc.tables, start=1):
        _extract_table_recursive(table, idx, out)
    return out


def _extract_textboxes(doc: Document, page_or_slide: int) -> List[TextBlock]:
    texts: List[TextBlock] = []
    try:
        nodes = doc.element.xpath(".//w:txbxContent//w:t")
        chunks = [n.text for n in nodes if n.text and n.text.strip()]
        if chunks:
            content = "\n".join(chunks)
            texts.append(TextBlock(content=content, page_or_slide=page_or_slide, block_type="paragraph"))
    except Exception:
        return texts
    return texts


def _iter_image_rels(doc: Document) -> Iterable:
    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            yield rel
    for section in doc.sections:
        for area in [section.header, section.footer]:
            for rel in area.part.rels.values():
                if "image" in rel.reltype:
                    yield rel


def _extract_images(doc: Document, images_dir: Path) -> List[ImageBlock]:
    image_blocks: List[ImageBlock] = []
    dedupe_keys = set()
    index = 0

    for rel in _iter_image_rels(doc):
        try:
            blob = rel.target_part.blob
        except Exception:
            continue
        blob_hash = hashlib.sha1(blob).hexdigest()
        if blob_hash in dedupe_keys:
            continue

        try:
            pil = Image.open(io.BytesIO(blob)).convert("RGB")
        except Exception:
            continue

        dedupe_key = dedupe_key_from_pil(pil)
        if dedupe_key in dedupe_keys:
            continue
        dedupe_keys.add(dedupe_key)
        dedupe_keys.add(blob_hash)

        index += 1
        out_name = f"doc_img{index}.png"
        out_path = images_dir / out_name
        pil.save(out_path)

        image_blocks.append(
            ImageBlock(
                image_path=str(out_path),
                page_or_slide=index,
                bounding_box=None,
                dimensions={"width": pil.width, "height": pil.height},
                extraction_method="embedded",
            )
        )

    return image_blocks


def extract_docx(file_path: str | Path, output_dir: str | Path) -> ExtractionResult:
    logger = setup_logger()
    file_path = Path(file_path)
    base_output = ensure_dir(output_dir)
    images_dir = ensure_dir(base_output / "images")

    try:
        doc = Document(str(file_path))
    except Exception as exc:
        raise RuntimeError(f"Failed to open DOCX: {file_path}") from exc

    warnings: List[str] = []
    texts: List[TextBlock] = []
    tables: List[TableBlock] = []
    images: List[ImageBlock] = []

    try:
        texts.extend(_extract_paragraph_texts(doc))
        texts.extend(_extract_header_footer_texts(doc, start_index=max(len(texts), 1)))
        texts.extend(_extract_textboxes(doc, page_or_slide=max(len(texts), 1)))
    except Exception as exc:
        warnings.append(f"Text extraction failed: {exc}")
        logger.warning(warnings[-1])

    try:
        for table in tqdm(_extract_tables(doc), desc=f"Extracting DOCX tables {file_path.name}"):
            tables.append(table)
    except Exception as exc:
        warnings.append(f"Table extraction failed: {exc}")
        logger.warning(warnings[-1])

    try:
        images.extend(_extract_images(doc, images_dir))
    except Exception as exc:
        warnings.append(f"Image extraction failed: {exc}")
        logger.warning(warnings[-1])

    return ExtractionResult(
        file_name=file_path.name,
        file_type="docx",
        extraction_timestamp=datetime.now(timezone.utc),
        num_pages_or_slides=max(len(texts), 1),
        texts=texts,
        tables=tables,
        images=images,
        warnings=warnings,
    )
