from __future__ import annotations

import io
import os
import tempfile
from pathlib import Path

import pytest


def _skip_if_missing(module_name: str) -> None:
    pytest.importorskip(module_name)


def _make_png_bytes() -> bytes:
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (220, 120), color=(255, 255, 255))
    draw = ImageDraw.Draw(image)
    draw.rectangle((10, 10, 210, 110), outline="black", width=2)
    draw.text((18, 45), "Demo Diagram", fill="black")
    buf = io.BytesIO()
    image.save(buf, format="PNG")
    return buf.getvalue()


def _create_sample_pdf(path: Path) -> None:
    import fitz

    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Quarterly Report")
    page.insert_text((72, 100), "ColA    ColB    ColC")
    page.insert_text((72, 120), "1       2       3")
    page.insert_text((72, 140), "4       5       6")

    rect = fitz.Rect(72, 180, 280, 300)
    page.insert_image(rect, stream=_make_png_bytes())

    # Vector drawing to trigger diagram rendering threshold.
    for y in range(330, 460, 6):
        page.draw_line((72, y), (500, y), color=(0, 0, 0), width=1)
    for x in range(72, 500, 12):
        page.draw_line((x, 330), (x, 460), color=(0, 0, 0), width=1)

    pdf.save(path)
    pdf.close()


def _create_sample_docx(path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading("System Manual", level=1)
    doc.add_paragraph("This section describes the process.")
    outer = doc.add_table(rows=2, cols=2)
    outer.cell(0, 0).text = "Header A"
    outer.cell(0, 1).text = "Header B"
    outer.cell(1, 0).text = "Value A1"
    nested = outer.cell(1, 1).add_table(rows=1, cols=2)
    nested.cell(0, 0).text = "Nested 1"
    nested.cell(0, 1).text = "Nested 2"

    img_path = path.with_suffix(".png")
    img_path.write_bytes(_make_png_bytes())
    doc.add_picture(str(img_path))
    doc.save(path)


def _create_sample_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
    title_box.text_frame.text = "Motor Controller"

    img_path = path.with_suffix(".png")
    img_path.write_bytes(_make_png_bytes())
    slide.shapes.add_picture(str(img_path), Inches(0.5), Inches(1.2), width=Inches(3.4))

    chart_data = CategoryChartData()
    chart_data.categories = ["Q1", "Q2", "Q3"]
    chart_data.add_series("Load", (5, 6, 7))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(4.2), Inches(1.2), Inches(5), Inches(3), chart_data)

    notes = slide.notes_slide
    notes.notes_text_frame.text = "Speaker notes for slide 1."

    prs.save(path)


def _create_sample_xlsx(path: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Summary"
    ws1["A1"] = "Metric"
    ws1["B1"] = "Value"
    ws1["A2"] = "Revenue"
    ws1["B2"] = "=Data!B2"
    ws1["A3"] = "Cost"
    ws1["B3"] = "=Data!C2"

    ws2 = wb.create_sheet("Data")
    ws2["A1"] = "Month"
    ws2["B1"] = "Revenue"
    ws2["C1"] = "Cost"
    ws2["A2"] = "Jan"
    ws2["B2"] = 100
    ws2["C2"] = 40
    ws2.merge_cells("A4:C4")
    ws2["A4"] = "Merged Footer"

    ws3 = wb.create_sheet("Calc")
    ws3["A1"] = "Profit"
    ws3["B1"] = "=Data!B2-Data!C2"

    wb.save(path)


def test_pdf_extractor_smoke():
    _skip_if_missing("fitz")
    _skip_if_missing("pdfplumber")
    from extractors.pdf_extractor import extract_pdf

    with tempfile.TemporaryDirectory() as tmp:
        base = Path(tmp)
        pdf_path = base / "sample.pdf"
        out = base / "out_pdf"
        _create_sample_pdf(pdf_path)
        result = extract_pdf(pdf_path, out)

        assert result.file_type == "pdf"
        assert result.num_pages_or_slides >= 1
        assert len(result.texts) > 0
        assert len(result.images) > 0


def test_docx_extractor_smoke():
    _skip_if_missing("docx")
    from extractors.docx_extractor import extract_docx

    with tempfile.TemporaryDirectory() as tmp:
        base = Path(tmp)
        docx_path = base / "sample.docx"
        out = base / "out_docx"
        _create_sample_docx(docx_path)
        result = extract_docx(docx_path, out)

        assert result.file_type == "docx"
        assert len(result.texts) > 0
        assert len(result.tables) > 0
        assert len(result.images) > 0


def test_pptx_extractor_smoke():
    _skip_if_missing("pptx")
    from extractors.pptx_extractor import extract_pptx

    with tempfile.TemporaryDirectory() as tmp:
        base = Path(tmp)
        pptx_path = base / "sample.pptx"
        out = base / "out_pptx"
        _create_sample_pptx(pptx_path)
        result = extract_pptx(pptx_path, out)

        assert result.file_type == "pptx"
        assert len(result.texts) > 0
        assert len(result.images) > 0


def test_spreadsheet_extractor_smoke():
    _skip_if_missing("openpyxl")
    _skip_if_missing("pandas")
    from extractors.spreadsheet_extractor import extract_spreadsheet

    with tempfile.TemporaryDirectory() as tmp:
        base = Path(tmp)
        xlsx_path = base / "sample.xlsx"
        out = base / "out_xlsx"
        _create_sample_xlsx(xlsx_path)
        result = extract_spreadsheet(xlsx_path, out)

        assert result.file_type == "xlsx"
        assert len(result.tables) >= 3
        assert result.sheet_dependencies is not None
        assert any(dep.references for dep in result.sheet_dependencies)
